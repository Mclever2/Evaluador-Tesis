"""Tests de la extracción DOCX (python-docx omite w:sdt en Paragraph.text)."""

from __future__ import annotations

from graph.segmenter import segmentar
from ingest.extractors import _sanear, extraer_documento
from rubrics import load_rubric

Z = chr(0x200B)  # ZERO WIDTH SPACE


def _docx_con_control_de_contenido(ruta):
    """DOCX con un control de contenido (w:sdt) en medio de un párrafo, como
    los campos de la carátula UPAO ('Autores:', 'TRUJILLO', etc.)."""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    parrafo = d.add_paragraph()
    sdt = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent>"
        "<w:r><w:t>Autores</w:t></w:r>"
        "</w:sdtContent></w:sdt>"
    )
    parrafo._p.append(sdt)
    parrafo.add_run(":")
    d.add_paragraph("Ponce Vásquez, McBreck")

    tabla = d.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text = "Tipo de investigación:"
    celda_sdt = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent>"
        "<w:r><w:t>Aplicada</w:t></w:r>"
        "</w:sdtContent></w:sdt>"
    )
    tabla.cell(0, 1).paragraphs[0]._p.append(celda_sdt)

    d.save(str(ruta))
    return ruta


class TestExtractorDocx:
    def test_incluye_texto_de_controles_de_contenido(self, tmp_path):
        """`Paragraph.text` omite w:sdt: se perdía 'Autores:' de la carátula y
        con ello la detección de nombres del anonimizador (caso real)."""
        ruta = _docx_con_control_de_contenido(tmp_path / "sdt.docx")
        doc = extraer_documento(ruta)
        texto = doc.texto_completo
        assert "Autores:" in texto
        assert "Ponce Vásquez, McBreck" in texto

    def test_incluye_sdt_dentro_de_celdas_de_tabla(self, tmp_path):
        ruta = _docx_con_control_de_contenido(tmp_path / "sdt2.docx")
        doc = extraer_documento(ruta)
        assert "Aplicada" in doc.texto_completo


class TestSaneadoAnchoCero:
    """Caso real ('PT_Lavado Flores'): la numeración de listas de Google Docs/Word
    arrastra un ZERO WIDTH SPACE tras el número ('1.​ Título', '1.3.​Importancia').
    Rompe la detección de numeración y el emparejamiento de aliases del
    segmentador; la extracción debe normalizarlo."""

    def test_sanear_sustituye_zwsp_por_espacio(self):
        # Ocupa el lugar de un espacio: '1.3.​Importancia' → '1.3. Importancia'
        assert _sanear(f"1.3.{Z}Importancia") == "1.3. Importancia"
        assert _sanear(f"1.{Z} Título") == "1.  Título"

    def test_sanear_elimina_juntadores_y_guion_suave(self):
        assert _sanear("cali" + chr(0x00AD) + "dad") == "calidad"  # SOFT HYPHEN
        assert _sanear("a" + chr(0xFEFF) + "b" + chr(0x2060) + "c") == "abc"  # BOM + WORD JOINER

    @staticmethod
    def _docx_con_numeracion_pegada(ruta):
        """DOCX cuyos encabezados numerados llevan el ancho cero pegado, como el
        PDF real. Sin saneado, S05 e S06 quedaban sin detectar."""
        import docx

        d = docx.Document()
        d.add_paragraph(f"1.{Z} PLANTEAMIENTO DEL ESTUDIO")
        d.add_paragraph(f"1.1.{Z} Descripción y delimitación del problema")
        d.add_paragraph(" ".join(f"p{i}" for i in range(80)))
        d.add_paragraph(f"1.3.{Z}Importancia del estudio")
        d.add_paragraph(" ".join(f"q{i}" for i in range(60)))
        d.add_paragraph(f"1.5.{Z}Limitaciones del estudio")
        d.add_paragraph(" ".join(f"r{i}" for i in range(40)))
        d.save(str(ruta))
        return ruta

    def test_extraccion_normaliza_el_ancho_cero(self, tmp_path):
        ruta = self._docx_con_numeracion_pegada(tmp_path / "zwsp.docx")
        texto = extraer_documento(ruta).texto_completo
        assert Z not in texto
        assert "1.3. Importancia del estudio" in texto

    def test_secciones_pegadas_al_numero_se_detectan(self, tmp_path):
        ruta = self._docx_con_numeracion_pegada(tmp_path / "zwsp2.docx")
        doc = extraer_documento(ruta)
        res = segmentar(doc, load_rubric("especifica_v1"))
        presentes = {s.seccion_id for s in res.presentes}
        assert {"S05", "S06"} <= presentes  # invisibles antes las ocultaban
        s05 = next(s for s in res.secciones if s.seccion_id == "S05")
        assert s05.encontrado_como == "Importancia del estudio"
